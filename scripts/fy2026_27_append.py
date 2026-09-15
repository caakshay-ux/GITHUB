#!/usr/bin/env python3
"""Shared helpers to append FY2026-27 YTD Capital Gains + Schedule OS sheets."""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from prepare_schedule_fa_ray import (
    FX_TO_USD_YE2025,
    SBI_TT_EUR,
    autosize,
    build_fifo_sales,
    build_lots_as_of,
    extract_dividends,
    extract_interest,
    extract_open_positions,
    extract_orders,
    extract_stock_splits,
    extract_transfers,
    extract_withholding,
    merge_orders,
    merge_splits,
    month_end_on_or_before,
    normalize_symbol,
    parse_ibkr,
    rule115_usd,
    style_header,
    to_usd,
    transfers_to_buy_orders,
    transfers_to_out_orders,
)


def sales_to_cg_rows(sales):
    rows = []
    for s in sales:
        if (s.code or "").upper() in ("I",):
            continue
        if abs(s.qty) < 1e-6 or (abs(s.proceeds_local) < 1e-8 and abs(s.cost_local) < 1e-8):
            continue
        if s.currency == "USD":
            proc_usd, cost_usd = s.proceeds_local, s.cost_local
        else:
            proc_usd = to_usd(s.proceeds_local, s.currency)
            cost_usd = to_usd(s.cost_local, s.currency)
        sale_rate_usd = rule115_usd(s.sell_date)
        cost_rate_usd = rule115_usd(s.acq_date)
        if s.currency == "EUR":
            sale_prev = (
                date(s.sell_date.year, s.sell_date.month, 1) - timedelta(days=1)
                if s.sell_date.month > 1
                else date(s.sell_date.year - 1, 12, 31)
            )
            cost_prev = (
                date(s.acq_date.year, s.acq_date.month, 1) - timedelta(days=1)
                if s.acq_date.month > 1
                else date(s.acq_date.year - 1, 12, 31)
            )
            sale_rate = SBI_TT_EUR[month_end_on_or_before(sale_prev, SBI_TT_EUR)]
            cost_rate = SBI_TT_EUR[month_end_on_or_before(cost_prev, SBI_TT_EUR)]
            sale_inr = s.proceeds_local * sale_rate
            cost_inr = s.cost_local * cost_rate
            comm_inr = abs(s.comm_usd) * sale_rate_usd
        else:
            sale_rate, cost_rate = sale_rate_usd, cost_rate_usd
            sale_inr = proc_usd * sale_rate_usd
            cost_inr = cost_usd * cost_rate_usd
            comm_inr = abs(s.comm_usd) * sale_rate_usd
        rows.append(
            {
                "symbol": s.symbol,
                "currency": s.currency,
                "qty": round(s.qty, 6),
                "acq_date": s.acq_date,
                "sell_date": s.sell_date,
                "holding_days": s.holding_days,
                "type": "LTCG" if s.is_ltcg else "STCG",
                "proceeds_local": round(s.proceeds_local, 2),
                "cost_local": round(s.cost_local, 2),
                "proceeds_usd": round(proc_usd, 2),
                "cost_usd": round(cost_usd, 2),
                "comm_usd": round(abs(s.comm_usd), 2),
                "sale_rate": sale_rate,
                "cost_rate": cost_rate,
                "sale_inr": round(sale_inr),
                "comm_inr": round(comm_inr),
                "cost_inr": round(cost_inr),
                "gain_inr": round(sale_inr - comm_inr - cost_inr),
                "note": "",
                "fcy_to_usd": 1.0 if s.currency in ("USD", "EUR") else FX_TO_USD_YE2025.get(s.currency, 1.0),
            }
        )
    return rows


def write_formula_cg(wb, title, subtitle, rows):
    from openpyxl.styles import Font, PatternFill

    ws = wb.create_sheet(title)
    ws["A1"] = subtitle
    ws.merge_cells("A1:U1")
    ws["A2"] = (
        "RECONCILIATION SHEET (formula-driven). Gain/(Loss) INR = Sale INR − Comm INR − Cost INR. "
        "Holding >730 → LTCG."
    )
    ws.merge_cells("A2:U2")
    ws["A3"] = "F=E-D | G=IF(F>730,\"LTCG\",\"STCG\") | L=H*K | M=I*K | Q=L*N | R=J*P | S=M*O | T=Q-R-S"
    ws.merge_cells("A3:U3")
    headers = [
        "Symbol", "CCY", "Qty", "Acquisition Date", "Sale Date", "Holding Days", "Type",
        "Sale Proceeds (FCY)", "Cost Basis (FCY)", "Comm (USD)", "FCY→USD Rate",
        "Sale Proceeds (USD)", "Cost (USD)", "Sale FX Rate", "Cost FX Rate", "Comm USD/INR Rate",
        "Sale (INR)", "Comm (INR)", "Cost (INR)", "Gain/(Loss) INR", "Notes",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header(ws, 4, len(headers))
    fill = PatternFill("solid", fgColor="FFF2CC")
    first = 5
    for i, r in enumerate(rows):
        rr = first + i
        usd_sale = rule115_usd(r["sell_date"])
        if r["currency"] == "EUR":
            fcy_to_usd, sale_fx, cost_fx = 1.0, r["sale_rate"], r["cost_rate"]
        else:
            fcy_to_usd = (
                1.0
                if r["currency"] == "USD"
                else r.get("fcy_to_usd", FX_TO_USD_YE2025.get(r["currency"], 1.0))
            )
            sale_fx = r["sale_rate"] if r["currency"] == "USD" else rule115_usd(r["sell_date"])
            cost_fx = r["cost_rate"] if r["currency"] == "USD" else rule115_usd(r["acq_date"])
        ws.cell(rr, 1, r["symbol"])
        ws.cell(rr, 2, r["currency"])
        ws.cell(rr, 3, r["qty"])
        ws.cell(rr, 4, r["acq_date"])
        ws.cell(rr, 5, r["sell_date"])
        ws.cell(rr, 6, f"=E{rr}-D{rr}")
        ws.cell(rr, 7, f'=IF(F{rr}>730,"LTCG","STCG")')
        ws.cell(rr, 8, round(r["proceeds_local"], 6))
        ws.cell(rr, 9, round(r["cost_local"], 6))
        ws.cell(rr, 10, r["comm_usd"])
        ws.cell(rr, 11, fcy_to_usd)
        ws.cell(rr, 12, f"=H{rr}*K{rr}")
        ws.cell(rr, 13, f"=I{rr}*K{rr}")
        ws.cell(rr, 14, sale_fx)
        ws.cell(rr, 15, cost_fx)
        ws.cell(rr, 16, usd_sale)
        ws.cell(rr, 17, f"=L{rr}*N{rr}")
        ws.cell(rr, 18, f"=J{rr}*P{rr}")
        ws.cell(rr, 19, f"=M{rr}*O{rr}")
        ws.cell(rr, 20, f"=Q{rr}-R{rr}-S{rr}")
        ws.cell(rr, 21, r.get("note") or "")
        for col in (8, 9, 10, 11, 14, 15, 16):
            ws.cell(rr, col).fill = fill
    last = first + len(rows) - 1 if rows else first - 1
    tot = last + 2
    if rows:
        ws.cell(tot, 7, "TOTAL STCG (INR)")
        ws.cell(tot, 20, f'=SUMIF(G{first}:G{last},"STCG",T{first}:T{last})')
        ws.cell(tot + 1, 7, "TOTAL LTCG (INR)")
        ws.cell(tot + 1, 20, f'=SUMIF(G{first}:G{last},"LTCG",T{first}:T{last})')
        ws.cell(tot + 2, 7, "TOTAL CG (INR)")
        ws.cell(tot + 2, 20, f"=T{tot}+T{tot+1}")
        for r in range(tot, tot + 3):
            ws.cell(r, 7).font = Font(bold=True)
            ws.cell(r, 20).font = Font(bold=True)
    autosize(ws)


def write_os_fy(wb, title, subtitle, dividends, interest):
    ws = wb.create_sheet(title)
    ws["A1"] = subtitle
    ws.append(
        ["Date", "Description", "Currency", "Amount (FCY)", "Amount (USD)", "USD/INR (Rule 115)", "Amount (INR)"]
    )
    style_header(ws, 2, 7)
    int_usd = int_inr = 0.0
    ws.append(["— INTEREST —"])
    for d in interest:
        usd = d["amount"] if d["currency"] == "USD" else to_usd(d["amount"], d["currency"])
        rate = rule115_usd(d["date"])
        inr = usd * rate
        int_usd += usd
        int_inr += inr
        ws.append([d["date"], d["description"], d["currency"], d["amount"], round(usd, 4), rate, round(inr)])
    ws.append([None, "TOTAL INTEREST", None, None, round(int_usd, 2), None, round(int_inr)])
    ws.append([])
    div_usd = div_inr = 0.0
    ws.append(["— DIVIDENDS —"])
    for d in dividends:
        usd = d["amount"] if d["currency"] == "USD" else to_usd(d["amount"], d["currency"])
        rate = rule115_usd(d["date"])
        inr = usd * rate
        div_usd += usd
        div_inr += inr
        ws.append([d["date"], d["description"], d["currency"], d["amount"], round(usd, 4), rate, round(inr)])
    ws.append([None, "TOTAL DIVIDENDS", None, None, round(div_usd, 2), None, round(div_inr)])
    autosize(ws)
    return div_usd, div_inr, int_usd, int_inr


def append_fy2026_27(
    out: Path,
    *,
    statement_paths: list[Path],
    fy_statement: Path,
    end: date,
    extra_notes: list[tuple[str, str]] | None = None,
) -> dict:
    """
    Append FY2026-27 YTD CG + OS using opening lots as of 31-Mar-2026 built from
    statement_paths (inception/fiscal/etc.) + fy_statement, then only in-window orders.
    """
    from openpyxl import load_workbook

    parsed = [parse_ibkr(p)[1] for p in statement_paths]
    sN = parse_ibkr(fy_statement)[1]
    open_snaps = [extract_open_positions(s) for s in parsed] + [extract_open_positions(sN)]
    all_xfers = []
    for s in parsed:
        all_xfers.extend(extract_transfers(s))
    all_xfers.extend(extract_transfers(sN))

    order_lists = [extract_orders(s) for s in parsed] + [extract_orders(sN)]
    orders = merge_orders(
        *order_lists,
        transfers_to_buy_orders(all_xfers, open_snaps),
        transfers_to_out_orders(all_xfers),
    )
    for o in orders:
        o["symbol"] = normalize_symbol(o["symbol"])
    splits = merge_splits(*[extract_stock_splits(s) for s in parsed], extract_stock_splits(sN))

    lots_open = build_lots_as_of(orders, splits, date(2026, 3, 31))
    orders_fy = [o for o in orders if o["date"] >= date(2026, 4, 1)]
    sales, _ = build_fifo_sales(lots_open, orders_fy, date(2026, 4, 1), end, splits=splits)
    cg_rows = sales_to_cg_rows(sales)
    stcg = sum(r["gain_inr"] for r in cg_rows if r["type"] == "STCG")
    ltcg = sum(r["gain_inr"] for r in cg_rows if r["type"] == "LTCG")

    divs = [d for d in extract_dividends(sN) if date(2026, 4, 1) <= d["date"] <= end]
    ints = [d for d in extract_interest(sN) if date(2026, 4, 1) <= d["date"] <= end]
    wht = [w for w in extract_withholding(sN) if date(2026, 4, 1) <= w["date"] <= end]

    end_label = end.strftime("%d-%b-%Y")
    wb = load_workbook(out)
    for name in ("Capital Gains FY2026-27", "Schedule OS - FY2026-27"):
        if name in wb.sheetnames:
            del wb[name]

    write_formula_cg(
        wb,
        "Capital Gains FY2026-27",
        f"Capital gains — FY 2026-27 YTD (01-Apr-2026 to {end_label}). "
        f"Partial year — statement ends {end_label}. FIFO. Yellow cells editable.",
        cg_rows,
    )
    div_usd, div_inr, int_usd, int_inr = write_os_fy(
        wb,
        "Schedule OS - FY2026-27",
        f"Interest & Dividend income — FY 2026-27 YTD (01-Apr-2026 to {end_label}). "
        "Partial — full year needs statement through 31-Mar-2027.",
        divs,
        ints,
    )

    if "FTC - Withholding Tax" in wb.sheetnames:
        wsw = wb["FTC - Withholding Tax"]
        for w in wht:
            usd = w["amount"] if w["currency"] == "USD" else to_usd(w["amount"], w["currency"])
            rate = rule115_usd(w["date"])
            wsw.append(
                [
                    "FY2026-27 YTD",
                    w["date"],
                    w["description"],
                    w["currency"],
                    w["amount"],
                    round(usd, 4),
                    rate,
                    round(usd * rate),
                ]
            )

    if "Notes for CA" in wb.sheetnames:
        wn = wb["Notes for CA"]
        wn.append([])
        wn.append([f"FY 2026-27 YTD (01-Apr-2026 to {end_label}) — partial statement"])
        wn.append(["Dividends FY2026-27 YTD", f"USD {div_usd:,.2f} / INR {div_inr:,.0f}"])
        wn.append(["Interest FY2026-27 YTD", f"USD {int_usd:,.2f} / INR {int_inr:,.0f}"])
        wn.append(["STCG FY2026-27 YTD (INR)", f"{stcg:,.0f}"])
        wn.append(["LTCG FY2026-27 YTD (INR)", f"{ltcg:,.0f}"])
        wn.append(
            [
                "FY2026-27 coverage",
                f"Activity Statement ends {end_label} — not full FY. "
                "Obtain statement to 31-Mar-2027 for complete year.",
            ]
        )
        for label, text in extra_notes or []:
            wn.append([label, text])

    if "CG Summary by Symbol" in wb.sheetnames:
        wss = wb["CG Summary by Symbol"]
        wss.append([])
        wss.append([f"— FY 2026-27 YTD (to {end_label}) —"])
        wss.append(["Symbol", "STCG INR", "LTCG INR", "Total Gain/(Loss) INR", "Sale Proceeds USD", "Cost USD"])
        by = {}
        for r in cg_rows:
            e = by.setdefault(r["symbol"], {"stcg": 0, "ltcg": 0, "sale": 0.0, "cost": 0.0})
            if r["type"] == "STCG":
                e["stcg"] += r["gain_inr"]
            else:
                e["ltcg"] += r["gain_inr"]
            e["sale"] += r["proceeds_usd"]
            e["cost"] += r["cost_usd"]
        for sym in sorted(by):
            e = by[sym]
            wss.append([sym, e["stcg"], e["ltcg"], e["stcg"] + e["ltcg"], round(e["sale"], 2), round(e["cost"], 2)])

    wb.save(out)
    print(
        f"FY2026-27 YTD: CG lots={len(cg_rows)} STCG={stcg:,.0f} LTCG={ltcg:,.0f} "
        f"Div USD={div_usd:,.2f} INR={div_inr:,.0f} Int USD={int_usd:,.2f}"
    )
    return {
        "stcg": stcg,
        "ltcg": ltcg,
        "div_usd": div_usd,
        "div_inr": div_inr,
        "int_usd": int_usd,
        "int_inr": int_inr,
        "lots": len(cg_rows),
    }
