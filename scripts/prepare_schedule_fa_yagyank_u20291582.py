#!/usr/bin/env python3
"""
Schedule FA / CG / Dividends for Yagyank Chadha — IBKR U20291582 (Individual).

Inputs: real IBKR Activity Statement FY2025-26 (23-May-2025–31-Mar-2026)
  + FY2026-27 YTD (01-Apr-2026–10-Sep-2026).
CY2025 Annual synthesized; YE marks from PortfolioAnalyst CY2025 Open Position Summary.
FY2025-26: inbound FOP/Internal only — nil taxable CG.
FY2026-27: taxable GOOG + SHOP sales.
Linked: Internal In of SHOP + Lux UCITS from U15172057 on 23-May-2025.
A3 Col C = Symbol only.
"""

from __future__ import annotations

import csv
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from prepare_schedule_fa_ray import (  # noqa: E402
    FX_TO_USD_YE2025,
    SBI_TT_EUR,
    autosize,
    build_fifo_sales,
    build_lots_as_of,
    extract_dividends,
    extract_interest,
    extract_open_positions,
    extract_orders,
    extract_stock_splits,
    extract_transfers,
    extract_withholding,
    merge_orders,
    merge_splits,
    month_end_on_or_before,
    normalize_symbol,
    parse_ibkr,
    rule115_usd,
    style_header,
    to_usd,
    transfers_to_buy_orders,
    transfers_to_out_orders,
    main,
)

INCEPTION = ROOT / "input_Inception_FY2025-26_Yagyank_U20291582.csv"
FISCAL = ROOT / "input_Fiscal_Statement_Yagyank_U20291582.csv"
FISCAL_FY2627 = ROOT / "input_Fiscal_Statement_FY2026-27_Yagyank_U20291582.csv"
ANNUAL = ROOT / "input_Annual_Statement_Yagyank_U20291582.csv"
OUT = ROOT / "output" / "Foreign_Assets_Schedule_FA_Yagyank_Chadha_U20291582_AY2026-27.xlsx"
FY2627_END = date(2026, 9, 10)

# YE2025 closes from PortfolioAnalyst CY2025 (same account)
PA_YE_CLOSE = {
    "LU0099574567": (80.78, "EUR"),
    "LU0106831901": (77.41, "USD"),
    "GOOG": (313.8, "USD"),
    "SHOP": (160.97, "USD"),
}


def _in_cy2025(d: date) -> bool:
    return date(2025, 1, 1) <= d <= date(2025, 12, 31)


def synthesize_annual(src: Path, annual_out: Path) -> None:
    _, s = parse_ibkr(src)
    open_pos = extract_open_positions(s)
    orders = merge_orders(
        extract_orders(s),
        transfers_to_buy_orders(extract_transfers(s), [open_pos]),
        transfers_to_out_orders(extract_transfers(s)),
    )
    for o in orders:
        o["symbol"] = normalize_symbol(o["symbol"])

    splits = merge_splits(extract_stock_splits(s))
    lots_ye = build_lots_as_of(orders, splits, date(2025, 12, 31))

    rows: list[list] = []

    def add(sec, kind, *vals):
        rows.append([sec, kind, *vals])

    add("Statement", "Header", "Field Name", "Field Value")
    add("Statement", "Data", "Title", "Activity Statement")
    add("Statement", "Data", "Period", "January 1, 2025 - December 31, 2025")
    add(
        "Statement",
        "Data",
        "WhenGenerated",
        "Synthesized for Schedule FA from FY Activity Statement; YE marks from PortfolioAnalyst CY2025",
    )

    add("Account Information", "Header", "Field Name", "Field Value")
    for _, r in s["Account Information"]:
        add("Account Information", "Data", *r)

    add(
        "Open Positions",
        "Header",
        "DataDiscriminator",
        "Asset Category",
        "Currency",
        "Symbol",
        "Quantity",
        "Mult",
        "Cost Price",
        "Cost Basis",
        "Close Price",
        "Value",
        "Unrealized P/L",
        "Code",
    )
    stock_usd = 0.0
    mtm_rows = []
    for sym, lots in sorted(lots_ye.items()):
        qty = sum(L.qty for L in lots)
        if qty <= 1e-8:
            continue
        cost = sum(L.cost_local for L in lots)
        ccy = lots[0].currency or "USD"
        px = (cost / qty) if qty else 0.0
        if sym in PA_YE_CLOSE and PA_YE_CLOSE[sym][1] == ccy:
            px = PA_YE_CLOSE[sym][0]
        value = px * qty
        stock_usd += to_usd(value, ccy)
        add(
            "Open Positions",
            "Data",
            "Summary",
            "Stocks" if not sym.startswith("LU") else "Mutual Funds",
            ccy,
            sym,
            round(qty, 6),
            1,
            round(cost / qty, 9) if qty else 0,
            round(cost, 6),
            round(px, 6),
            round(value, 6),
            round(value - cost, 6),
            "",
        )
        mtm_rows.append((sym, qty, px))

    # Cash from FY ending cash is after Mar — use PA YE cash ~16.54
    cash = 16.54
    ending = round(stock_usd + cash, 2)

    add("Net Asset Value", "Header", "Asset Class", "Prior Total", "Current Long", "Current Short", "Current Total", "Change")
    add("Net Asset Value", "Data", "Cash ", 0, cash, 0, cash, cash)
    add("Net Asset Value", "Data", "Stock", 0, round(stock_usd, 2), 0, round(stock_usd, 2), round(stock_usd, 2))
    add("Net Asset Value", "Data", "Total", 0, ending, 0, ending, ending)
    add("Change in NAV", "Header", "Field Name", "Field Value")
    add("Change in NAV", "Data", "Starting Value", 0)
    add("Change in NAV", "Data", "Ending Value", ending)

    add(
        "Mark-to-Market Performance Summary",
        "Header",
        "Asset Category",
        "Symbol",
        "Prior Quantity",
        "Current Quantity",
        "Prior Price",
        "Current Price",
        "Mark-to-Market P/L Position",
        "Mark-to-Market P/L Transaction",
        "Mark-to-Market P/L Commissions",
        "Mark-to-Market P/L Other",
        "Mark-to-Market P/L Total",
        "Code",
    )
    for sym, qty, px in mtm_rows:
        add("Mark-to-Market Performance Summary", "Data", "Stocks", sym, 0, qty, "--", px, 0, 0, 0, 0, 0, "")

    add(
        "Transfers",
        "Header",
        "Asset Category",
        "Currency",
        "Symbol",
        "Date",
        "Type",
        "Direction",
        "Xfer Company",
        "Xfer Account",
        "Qty",
        "Xfer Price",
        "Market Value",
        "Realized P/L",
        "Cash Amount",
        "Code",
    )
    for kind, r in s.get("Transfers", []):
        if kind != "Data" or r[0] == "Total" or str(r[0]).startswith("Total"):
            continue
        d = datetime.strptime(str(r[3])[:10], "%Y-%m-%d").date()
        if _in_cy2025(d):
            add("Transfers", "Data", *r)

    add(
        "Trades",
        "Header",
        "DataDiscriminator",
        "Asset Category",
        "Currency",
        "Symbol",
        "Date/Time",
        "Quantity",
        "T. Price",
        "C. Price",
        "Proceeds",
        "Comm/Fee",
        "Basis",
        "Realized P/L",
        "MTM P/L",
        "Code",
    )

    add("Dividends", "Header", "Currency", "Date", "Description", "Amount")
    for d in sorted([x for x in extract_dividends(s) if _in_cy2025(x["date"])], key=lambda x: x["date"]):
        add("Dividends", "Data", d["currency"], d["date"].isoformat(), d["description"], d["amount"])

    add("Withholding Tax", "Header", "Currency", "Date", "Description", "Amount", "Code")
    for w in sorted([x for x in extract_withholding(s) if _in_cy2025(x["date"])], key=lambda x: x["date"]):
        add("Withholding Tax", "Data", w["currency"], w["date"].isoformat(), w["description"], w["amount"], "")

    add("Deposits & Withdrawals", "Header", "Currency", "Settle Date", "Description", "Amount")
    add("Interest", "Header", "Currency", "Date", "Description", "Amount")
    for d in sorted([x for x in extract_interest(s) if _in_cy2025(x["date"])], key=lambda x: x["date"]):
        add("Interest", "Data", d["currency"], d["date"].isoformat(), d["description"], d["amount"])

    add(
        "Financial Instrument Information",
        "Header",
        "Asset Category",
        "Symbol",
        "Description",
        "Conid",
        "Security ID",
        "Underlying",
        "Listing Exch",
        "Multiplier",
        "Type",
        "Code",
    )
    seen = set()
    for kind, r in s.get("Financial Instrument Information", []):
        if kind != "Data":
            continue
        key = r[1]
        if key in seen:
            continue
        seen.add(key)
        add("Financial Instrument Information", "Data", *r)

    annual_out.parent.mkdir(parents=True, exist_ok=True)
    with open(annual_out, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)
    print(f"Synthesized annual → {annual_out} ({len(rows)} rows, YE symbols={len(mtm_rows)}, NAV~{ending})")


NOTES = [
    (
        "Assessee",
        "Yagyank Chadha (also spelled Chaddha) — IBKR U20291582 (Individual) — "
        "Advisor: MATCAP WEALTH ADVISORS PRIVATE LIMITED. Satellite / Chaddha_1 account.",
    ),
    (
        "Source data",
        "Real IBKR Activity Statement FY2025-26 (23-May-2025–31-Mar-2026). "
        "FY2026-27 YTD Activity Statement (01-Apr-2026–10-Sep-2026). "
        "CY2025 Annual synthesized; YE marks from PortfolioAnalyst CY2025 Open Position Summary.",
    ),
    (
        "Account opening",
        "Opened 23-May-2025 via Internal In from U15172057: SHOP × 84, "
        "Fidelity Global Technology LU0099574567 × 511.42, BGF World Financials LU0106831901 × 792.45. "
        "FOP/ACATS In GOOG × 105 on 01-Dec-2025. FOP In GOOG × 323 on 08-Sep-2026.",
    ),
    (
        "Related account",
        "Assets transferred in from main trading account U15172057 (same beneficial owner). "
        "See U15172057 workbook for Internal Out (code I — not taxable CG).",
    ),
    (
        "Capital gains FY2025-26",
        "Nil taxable stock/fund sales in FY2025-26 (inbound transfers only).",
    ),
    (
        "Capital gains FY2026-27 YTD",
        "Partial (to 10-Sep-2026): taxable sales of GOOG × 105 (12-May-2026) and SHOP × 44 "
        "(Jul–Aug 2026). Acquisition dates = inbound transfer dates (confirm original buys "
        "at U15172057 / prior broker for LTCG — workbook uses transfer dates → STCG if <730 days).",
    ),
    (
        "Dividends / FTC",
        "FY2025-26: GOOG USD 22.05 on 15-Dec-2025 and 16-Mar-2026 (FY total USD 44.10); "
        "US withholding 25% (USD 5.51 × 2) — see FTC sheet. "
        "FY2026-27 YTD: nil dividends; interest only (USD credit interest).",
    ),
    (
        "Cost / acquisition dates",
        "Cost basis from IBKR Open Positions carryover. Acquisition dates = inbound transfer dates. "
        "Confirm original buy dates at U15172057 / prior broker for LTCG on future sales "
        "(SHOP/funds originally FOP into U15172057 in Sep-2024).",
    ),
]


def _sales_to_cg_rows(sales):
    rows = []
    for s in sales:
        if (s.code or "").upper() in ("I",):
            continue
        if abs(s.qty) < 1e-6 or (abs(s.proceeds_local) < 1e-8 and abs(s.cost_local) < 1e-8):
            continue
        if s.currency == "USD":
            proc_usd, cost_usd = s.proceeds_local, s.cost_local
        else:
            proc_usd = to_usd(s.proceeds_local, s.currency)
            cost_usd = to_usd(s.cost_local, s.currency)
        sale_rate_usd = rule115_usd(s.sell_date)
        cost_rate_usd = rule115_usd(s.acq_date)
        if s.currency == "EUR":
            sale_prev = (
                date(s.sell_date.year, s.sell_date.month, 1) - timedelta(days=1)
                if s.sell_date.month > 1
                else date(s.sell_date.year - 1, 12, 31)
            )
            cost_prev = (
                date(s.acq_date.year, s.acq_date.month, 1) - timedelta(days=1)
                if s.acq_date.month > 1
                else date(s.acq_date.year - 1, 12, 31)
            )
            sale_rate = SBI_TT_EUR[month_end_on_or_before(sale_prev, SBI_TT_EUR)]
            cost_rate = SBI_TT_EUR[month_end_on_or_before(cost_prev, SBI_TT_EUR)]
            sale_inr = s.proceeds_local * sale_rate
            cost_inr = s.cost_local * cost_rate
            comm_inr = abs(s.comm_usd) * sale_rate_usd
        else:
            sale_rate, cost_rate = sale_rate_usd, cost_rate_usd
            sale_inr = proc_usd * sale_rate_usd
            cost_inr = cost_usd * cost_rate_usd
            comm_inr = abs(s.comm_usd) * sale_rate_usd
        rows.append(
            {
                "symbol": s.symbol,
                "currency": s.currency,
                "qty": round(s.qty, 6),
                "acq_date": s.acq_date,
                "sell_date": s.sell_date,
                "holding_days": s.holding_days,
                "type": "LTCG" if s.is_ltcg else "STCG",
                "proceeds_local": round(s.proceeds_local, 2),
                "cost_local": round(s.cost_local, 2),
                "proceeds_usd": round(proc_usd, 2),
                "cost_usd": round(cost_usd, 2),
                "comm_usd": round(abs(s.comm_usd), 2),
                "sale_rate": sale_rate,
                "cost_rate": cost_rate,
                "sale_inr": round(sale_inr),
                "comm_inr": round(comm_inr),
                "cost_inr": round(cost_inr),
                "gain_inr": round(sale_inr - comm_inr - cost_inr),
                "note": "",
                "fcy_to_usd": 1.0 if s.currency in ("USD", "EUR") else FX_TO_USD_YE2025.get(s.currency, 1.0),
            }
        )
    return rows


def _write_formula_cg(wb, title, subtitle, rows):
    from openpyxl.styles import Font, PatternFill

    ws = wb.create_sheet(title)
    ws["A1"] = subtitle
    ws.merge_cells("A1:U1")
    ws["A2"] = "RECONCILIATION SHEET (formula-driven). Gain/(Loss) INR = Sale INR − Comm INR − Cost INR. Holding >730 → LTCG."
    ws.merge_cells("A2:U2")
    ws["A3"] = "F=E-D | G=IF(F>730,\"LTCG\",\"STCG\") | L=H*K | M=I*K | Q=L*N | R=J*P | S=M*O | T=Q-R-S"
    ws.merge_cells("A3:U3")
    headers = [
        "Symbol", "CCY", "Qty", "Acquisition Date", "Sale Date", "Holding Days", "Type",
        "Sale Proceeds (FCY)", "Cost Basis (FCY)", "Comm (USD)", "FCY→USD Rate",
        "Sale Proceeds (USD)", "Cost (USD)", "Sale FX Rate", "Cost FX Rate", "Comm USD/INR Rate",
        "Sale (INR)", "Comm (INR)", "Cost (INR)", "Gain/(Loss) INR", "Notes",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header(ws, 4, len(headers))
    fill = PatternFill("solid", fgColor="FFF2CC")
    first = 5
    for i, r in enumerate(rows):
        rr = first + i
        usd_sale = rule115_usd(r["sell_date"])
        if r["currency"] == "EUR":
            fcy_to_usd, sale_fx, cost_fx = 1.0, r["sale_rate"], r["cost_rate"]
        else:
            fcy_to_usd = 1.0 if r["currency"] == "USD" else r.get("fcy_to_usd", FX_TO_USD_YE2025.get(r["currency"], 1.0))
            sale_fx = r["sale_rate"] if r["currency"] == "USD" else rule115_usd(r["sell_date"])
            cost_fx = r["cost_rate"] if r["currency"] == "USD" else rule115_usd(r["acq_date"])
        ws.cell(rr, 1, r["symbol"])
        ws.cell(rr, 2, r["currency"])
        ws.cell(rr, 3, r["qty"])
        ws.cell(rr, 4, r["acq_date"])
        ws.cell(rr, 5, r["sell_date"])
        ws.cell(rr, 6, f"=E{rr}-D{rr}")
        ws.cell(rr, 7, f'=IF(F{rr}>730,"LTCG","STCG")')
        ws.cell(rr, 8, round(r["proceeds_local"], 6))
        ws.cell(rr, 9, round(r["cost_local"], 6))
        ws.cell(rr, 10, r["comm_usd"])
        ws.cell(rr, 11, fcy_to_usd)
        ws.cell(rr, 12, f"=H{rr}*K{rr}")
        ws.cell(rr, 13, f"=I{rr}*K{rr}")
        ws.cell(rr, 14, sale_fx)
        ws.cell(rr, 15, cost_fx)
        ws.cell(rr, 16, usd_sale)
        ws.cell(rr, 17, f"=L{rr}*N{rr}")
        ws.cell(rr, 18, f"=J{rr}*P{rr}")
        ws.cell(rr, 19, f"=M{rr}*O{rr}")
        ws.cell(rr, 20, f"=Q{rr}-R{rr}-S{rr}")
        ws.cell(rr, 21, r.get("note") or "")
        for col in (8, 9, 10, 11, 14, 15, 16):
            ws.cell(rr, col).fill = fill
    last = first + len(rows) - 1 if rows else first - 1
    tot = last + 2
    if rows:
        ws.cell(tot, 7, "TOTAL STCG (INR)")
        ws.cell(tot, 20, f'=SUMIF(G{first}:G{last},"STCG",T{first}:T{last})')
        ws.cell(tot + 1, 7, "TOTAL LTCG (INR)")
        ws.cell(tot + 1, 20, f'=SUMIF(G{first}:G{last},"LTCG",T{first}:T{last})')
        ws.cell(tot + 2, 7, "TOTAL CG (INR)")
        ws.cell(tot + 2, 20, f"=T{tot}+T{tot+1}")
        for r in range(tot, tot + 3):
            ws.cell(r, 7).font = Font(bold=True)
            ws.cell(r, 20).font = Font(bold=True)
    autosize(ws)


def _write_os_fy(wb, title, subtitle, dividends, interest):
    ws = wb.create_sheet(title)
    ws["A1"] = subtitle
    ws.append(["Date", "Description", "Currency", "Amount (FCY)", "Amount (USD)", "USD/INR (Rule 115)", "Amount (INR)"])
    style_header(ws, 2, 7)
    int_usd = int_inr = 0.0
    ws.append(["— INTEREST —"])
    for d in interest:
        usd = d["amount"] if d["currency"] == "USD" else to_usd(d["amount"], d["currency"])
        rate = rule115_usd(d["date"])
        inr = usd * rate
        int_usd += usd
        int_inr += inr
        ws.append([d["date"], d["description"], d["currency"], d["amount"], round(usd, 4), rate, round(inr)])
    ws.append([None, "TOTAL INTEREST", None, None, round(int_usd, 2), None, round(int_inr)])
    ws.append([])
    div_usd = div_inr = 0.0
    ws.append(["— DIVIDENDS —"])
    for d in dividends:
        usd = d["amount"] if d["currency"] == "USD" else to_usd(d["amount"], d["currency"])
        rate = rule115_usd(d["date"])
        inr = usd * rate
        div_usd += usd
        div_inr += inr
        ws.append([d["date"], d["description"], d["currency"], d["amount"], round(usd, 4), rate, round(inr)])
    ws.append([None, "TOTAL DIVIDENDS", None, None, round(div_usd, 2), None, round(div_inr)])
    autosize(ws)
    return div_usd, div_inr, int_usd, int_inr


def append_fy2026_27(out: Path) -> dict:
    """Add FY2026-27 YTD (to 10-Sep-2026) CG + OS sheets."""
    from openpyxl import load_workbook

    _, sI = parse_ibkr(INCEPTION)
    _, sF = parse_ibkr(FISCAL)
    _, sN = parse_ibkr(FISCAL_FY2627)
    open_inc = extract_open_positions(sI)
    open_fy = extract_open_positions(sF)
    open_n = extract_open_positions(sN)
    all_xfers = extract_transfers(sI) + extract_transfers(sF) + extract_transfers(sN)

    orders = merge_orders(
        extract_orders(sI),
        extract_orders(sF),
        extract_orders(sN),
        transfers_to_buy_orders(all_xfers, [open_inc, open_fy, open_n]),
        transfers_to_out_orders(all_xfers),
    )
    for o in orders:
        o["symbol"] = normalize_symbol(o["symbol"])
    splits = merge_splits(
        extract_stock_splits(sI),
        extract_stock_splits(sF),
        extract_stock_splits(sN),
    )

    lots_open = build_lots_as_of(orders, splits, date(2026, 3, 31))
    orders_fy = [o for o in orders if o["date"] >= date(2026, 4, 1)]
    end = FY2627_END
    sales, _ = build_fifo_sales(lots_open, orders_fy, date(2026, 4, 1), end, splits=splits)
    cg_rows = _sales_to_cg_rows(sales)
    stcg = sum(r["gain_inr"] for r in cg_rows if r["type"] == "STCG")
    ltcg = sum(r["gain_inr"] for r in cg_rows if r["type"] == "LTCG")

    divs = [d for d in extract_dividends(sN) if date(2026, 4, 1) <= d["date"] <= end]
    ints = [d for d in extract_interest(sN) if date(2026, 4, 1) <= d["date"] <= end]
    wht = [w for w in extract_withholding(sN) if date(2026, 4, 1) <= w["date"] <= end]

    wb = load_workbook(out)
    for name in ("Capital Gains FY2026-27", "Schedule OS - FY2026-27"):
        if name in wb.sheetnames:
            del wb[name]

    _write_formula_cg(
        wb,
        "Capital Gains FY2026-27",
        "Capital gains — FY 2026-27 YTD (01-Apr-2026 to 10-Sep-2026). Partial year — statement ends 10-Sep-2026. FIFO. Yellow cells editable.",
        cg_rows,
    )
    div_usd, div_inr, int_usd, int_inr = _write_os_fy(
        wb,
        "Schedule OS - FY2026-27",
        "Interest & Dividend income — FY 2026-27 YTD (01-Apr-2026 to 10-Sep-2026). Partial — full year needs statement through 31-Mar-2027.",
        divs,
        ints,
    )

    if "FTC - Withholding Tax" in wb.sheetnames:
        wsw = wb["FTC - Withholding Tax"]
        for w in wht:
            usd = w["amount"] if w["currency"] == "USD" else to_usd(w["amount"], w["currency"])
            rate = rule115_usd(w["date"])
            wsw.append(
                [
                    "FY2026-27 YTD",
                    w["date"],
                    w["description"],
                    w["currency"],
                    w["amount"],
                    round(usd, 4),
                    rate,
                    round(usd * rate),
                ]
            )

    if "Notes for CA" in wb.sheetnames:
        wn = wb["Notes for CA"]
        wn.append([])
        wn.append(["FY 2026-27 YTD (01-Apr-2026 to 10-Sep-2026) — partial statement"])
        wn.append(["Dividends FY2026-27 YTD", f"USD {div_usd:,.2f} / INR {div_inr:,.0f}"])
        wn.append(["Interest FY2026-27 YTD", f"USD {int_usd:,.2f} / INR {int_inr:,.0f}"])
        wn.append(["STCG FY2026-27 YTD (INR)", f"{stcg:,.0f}"])
        wn.append(["LTCG FY2026-27 YTD (INR)", f"{ltcg:,.0f}"])
        wn.append(
            [
                "FY2026-27 coverage",
                "Activity Statement ends 10-Sep-2026 — not full FY. Obtain statement to 31-Mar-2027 for complete year.",
            ]
        )
        wn.append(
            [
                "GOOG FOP 08-Sep-2026",
                "FOP In GOOG × 323 after May sale — held at period end; not taxable CG. "
                "Confirm original buy dates for LTCG on future sales.",
            ]
        )

    if "CG Summary by Symbol" in wb.sheetnames:
        wss = wb["CG Summary by Symbol"]
        wss.append([])
        wss.append(["— FY 2026-27 YTD (to 10-Sep-2026) —"])
        wss.append(["Symbol", "STCG INR", "LTCG INR", "Total Gain/(Loss) INR", "Sale Proceeds USD", "Cost USD"])
        by = {}
        for r in cg_rows:
            e = by.setdefault(r["symbol"], {"stcg": 0, "ltcg": 0, "sale": 0.0, "cost": 0.0})
            if r["type"] == "STCG":
                e["stcg"] += r["gain_inr"]
            else:
                e["ltcg"] += r["gain_inr"]
            e["sale"] += r["proceeds_usd"]
            e["cost"] += r["cost_usd"]
        for sym in sorted(by):
            e = by[sym]
            wss.append([sym, e["stcg"], e["ltcg"], e["stcg"] + e["ltcg"], round(e["sale"], 2), round(e["cost"], 2)])

    wb.save(out)
    print(
        f"FY2026-27 YTD: CG lots={len(cg_rows)} STCG={stcg:,.0f} LTCG={ltcg:,.0f} "
        f"Div USD={div_usd:,.2f} INR={div_inr:,.0f} Int USD={int_usd:,.2f}"
    )
    return {"stcg": stcg, "ltcg": ltcg, "div_usd": div_usd, "div_inr": div_inr, "int_usd": int_usd}


def run():
    synthesize_annual(FISCAL, ANNUAL)
    main(
        inception=INCEPTION,
        annual=ANNUAL,
        fiscal=FISCAL,
        out=OUT,
        account_open_date="2025-05-23 (Internal In from U15172057)",
        assessee_label="Yagyank Chadha (U20291582)",
        notes_extra=NOTES,
    )
    append_fy2026_27(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    run()
