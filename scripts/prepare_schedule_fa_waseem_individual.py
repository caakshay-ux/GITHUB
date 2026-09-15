#!/usr/bin/env python3
"""
Schedule FA / CG / Dividends for Mohammad Waseem — IBKR U25039405 (Individual).

Sources:
  - FY 2026-27 YTD Activity Statement (01-Apr-2026–03-Sep-2026) only
  - FY2025-26 / CY2025 stubs + synthesized Annual (1 GOOG held at YE2025)

GOOG: prior × 1 + FOP In × 124 (11-May-2026); Internal transfers with Joint U24577010
(code I — not taxable). Taxable sell GOOG × 25 on 01-Jul-2026.
A3 Col C = Symbol only.
"""

from __future__ import annotations

import csv
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from prepare_schedule_fa_ray import (  # noqa: E402
    FX_TO_USD_YE2025,
    SBI_TT_EUR,
    autosize,
    build_fifo_sales,
    build_lots_as_of,
    extract_dividends,
    extract_interest,
    extract_open_positions,
    extract_orders,
    extract_stock_splits,
    extract_transfers,
    extract_withholding,
    main,
    merge_orders,
    merge_splits,
    month_end_on_or_before,
    normalize_symbol,
    parse_ibkr,
    rule115_usd,
    style_header,
    to_usd,
    transfers_to_out_orders,
)

INCEPTION = ROOT / "input_Inception_FY2024-25_Waseem_Individual.csv"
FISCAL = ROOT / "input_Fiscal_Statement_Waseem_Individual.csv"
FISCAL_FY2627 = ROOT / "input_Fiscal_Statement_FY2026-27_Waseem_Individual.csv"
ANNUAL = ROOT / "input_Annual_Statement_Waseem_Individual.csv"
OUT = ROOT / "output" / "Foreign_Assets_Schedule_FA_Mohammad_Waseem_U25039405_AY2026-27.xlsx"

# IBKR sell Basis implies FOP unit 99.21; prior share unit 98.82 (FIFO into Joint first sell).
GOOG_PRIOR_QTY = 1.0
GOOG_PRIOR_COST = 98.82
GOOG_FOP_QTY = 124.0
GOOG_FOP_UNIT = 99.21
GOOG_FOP_COST = GOOG_FOP_QTY * GOOG_FOP_UNIT  # 12,302.04
GOOG_INTERNAL_IN_QTY = 28.0
GOOG_INTERNAL_IN_COST = GOOG_INTERNAL_IN_QTY * GOOG_FOP_UNIT


def _seed(sym, ccy, qty, cost, d: date, sec: int = 0) -> dict:
    return {
        "currency": ccy,
        "symbol": normalize_symbol(sym),
        "datetime": f"{d.isoformat()}, 00:00:{sec:02d}",
        "date": d,
        "qty": qty,
        "price": cost / qty if qty else 0.0,
        "proceeds": -cost,
        "comm_usd": 0.0,
        "basis": cost,
        "realized": 0.0,
        "code": "O",
    }


def seed_orders() -> list[dict]:
    return [
        _seed("GOOG", "USD", GOOG_PRIOR_QTY, GOOG_PRIOR_COST, date(2024, 12, 31), 0),
        _seed("GOOG", "USD", GOOG_FOP_QTY, GOOG_FOP_COST, date(2026, 5, 11), 1),
        _seed("GOOG", "USD", GOOG_INTERNAL_IN_QTY, GOOG_INTERNAL_IN_COST, date(2026, 8, 20), 2),
    ]


def synthesize_annual(annual_out: Path, seeds: list[dict]) -> None:
    """CY2025: only prior GOOG × 1 (account activity otherwise starts FY2026-27)."""
    _, sF = parse_ibkr(FISCAL_FY2627)
    lots = {}
    for o in seeds:
        if o["date"] <= date(2025, 12, 31) and o["qty"] > 0:
            lots.setdefault(o["symbol"], []).append(o)

    rows: list[list] = []

    def add(sec, kind, *vals):
        rows.append([sec, kind, *vals])

    add("Statement", "Header", "Field Name", "Field Value")
    add("Statement", "Data", "Title", "Activity Statement")
    add("Statement", "Data", "Period", "January 1, 2025 - December 31, 2025")
    add(
        "Statement",
        "Data",
        "WhenGenerated",
        "Synthesized — FY2026-27 YTD statement only; YE GOOG × 1 at seeded cost",
    )
    add("Account Information", "Header", "Field Name", "Field Value")
    for _, r in sF["Account Information"]:
        add("Account Information", "Data", *r)

    add(
        "Open Positions",
        "Header",
        "DataDiscriminator",
        "Asset Category",
        "Currency",
        "Symbol",
        "Quantity",
        "Mult",
        "Cost Price",
        "Cost Basis",
        "Close Price",
        "Value",
        "Unrealized P/L",
        "Code",
    )
    stock_usd = 0.0
    mtm_rows = []
    for sym, olots in sorted(lots.items()):
        qty = sum(o["qty"] for o in olots)
        cost = sum(o["basis"] for o in olots)
        ccy = olots[0]["currency"]
        px = cost / qty if qty else 0.0
        value = px * qty
        stock_usd += to_usd(value, ccy)
        add(
            "Open Positions",
            "Data",
            "Summary",
            "Stocks",
            ccy,
            sym,
            qty,
            1,
            round(px, 9),
            round(cost, 6),
            round(px, 6),
            round(value, 6),
            0,
            "",
        )
        mtm_rows.append((sym, qty, px))

    ending = round(stock_usd, 2)
    add("Net Asset Value", "Header", "Asset Class", "Prior Total", "Current Long", "Current Short", "Current Total", "Change")
    add("Net Asset Value", "Data", "Total", ending, ending, 0, ending, 0)
    add("Change in NAV", "Header", "Field Name", "Field Value")
    add("Change in NAV", "Data", "Starting Value", ending)
    add("Change in NAV", "Data", "Ending Value", ending)
    add(
        "Mark-to-Market Performance Summary",
        "Header",
        "Asset Category",
        "Symbol",
        "Prior Quantity",
        "Current Quantity",
        "Prior Price",
        "Current Price",
        "Mark-to-Market P/L Position",
        "Mark-to-Market P/L Transaction",
        "Mark-to-Market P/L Commissions",
        "Mark-to-Market P/L Other",
        "Mark-to-Market P/L Total",
        "Code",
    )
    for sym, qty, px in mtm_rows:
        add("Mark-to-Market Performance Summary", "Data", "Stocks", sym, qty, qty, px, px, 0, 0, 0, 0, 0, "")

    for sec, hdr in [
        ("Trades", ["DataDiscriminator", "Asset Category", "Currency", "Symbol", "Date/Time", "Quantity", "T. Price", "C. Price", "Proceeds", "Comm/Fee", "Basis", "Realized P/L", "MTM P/L", "Code"]),
        ("Dividends", ["Currency", "Date", "Description", "Amount"]),
        ("Interest", ["Currency", "Date", "Description", "Amount"]),
        ("Withholding Tax", ["Currency", "Date", "Description", "Amount", "Code"]),
        ("Transfers", ["Asset Category", "Currency", "Symbol", "Date", "Type", "Direction", "Xfer Company", "Xfer Account", "Qty", "Xfer Price", "Market Value", "Realized P/L", "Cash Amount", "Code"]),
        ("Corporate Actions", ["Asset Category", "Currency", "Report Date", "Date/Time", "Description", "Quantity", "Proceeds", "Value", "Realized P/L", "Code"]),
        ("Financial Instrument Information", ["Asset Category", "Symbol", "Description", "Conid", "Security ID", "Underlying", "Listing Exch", "Multiplier", "Type", "Code"]),
    ]:
        add(sec, "Header", *hdr)
    for kind, r in sF.get("Financial Instrument Information", []):
        if kind == "Data":
            add("Financial Instrument Information", "Data", *r)

    annual_out.parent.mkdir(parents=True, exist_ok=True)
    with open(annual_out, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)
    print(f"Synthesized annual → {annual_out} (YE NAV~{ending})")


def _sales_to_cg_rows(sales):
    rows = []
    for s in sales:
        if (s.code or "").upper() in ("I",):
            continue
        if abs(s.qty) < 1e-6 or (abs(s.proceeds_local) < 1e-8 and abs(s.cost_local) < 1e-8):
            continue
        if s.currency == "USD":
            proc_usd, cost_usd = s.proceeds_local, s.cost_local
        else:
            proc_usd = to_usd(s.proceeds_local, s.currency)
            cost_usd = to_usd(s.cost_local, s.currency)
        sale_rate_usd = rule115_usd(s.sell_date)
        cost_rate_usd = rule115_usd(s.acq_date)
        if s.currency == "EUR":
            sale_prev = date(s.sell_date.year, s.sell_date.month, 1) - timedelta(days=1) if s.sell_date.month > 1 else date(s.sell_date.year - 1, 12, 31)
            cost_prev = date(s.acq_date.year, s.acq_date.month, 1) - timedelta(days=1) if s.acq_date.month > 1 else date(s.acq_date.year - 1, 12, 31)
            sale_rate = SBI_TT_EUR[month_end_on_or_before(sale_prev, SBI_TT_EUR)]
            cost_rate = SBI_TT_EUR[month_end_on_or_before(cost_prev, SBI_TT_EUR)]
            sale_inr = s.proceeds_local * sale_rate
            cost_inr = s.cost_local * cost_rate
            comm_inr = abs(s.comm_usd) * sale_rate_usd
        else:
            sale_rate, cost_rate = sale_rate_usd, cost_rate_usd
            sale_inr = proc_usd * sale_rate_usd
            cost_inr = cost_usd * cost_rate_usd
            comm_inr = abs(s.comm_usd) * sale_rate_usd
        note = ""
        if s.symbol == "GOOG":
            note = "Confirm original FOP/prior-broker buy dates for LTCG (IBKR Realized tagged LT; workbook uses transfer/seed dates)"
        rows.append(
            {
                "symbol": s.symbol,
                "currency": s.currency,
                "qty": round(s.qty, 6),
                "acq_date": s.acq_date,
                "sell_date": s.sell_date,
                "holding_days": s.holding_days,
                "type": "LTCG" if s.is_ltcg else "STCG",
                "proceeds_local": round(s.proceeds_local, 2),
                "cost_local": round(s.cost_local, 2),
                "proceeds_usd": round(proc_usd, 2),
                "cost_usd": round(cost_usd, 2),
                "comm_usd": round(abs(s.comm_usd), 2),
                "sale_rate": sale_rate,
                "cost_rate": cost_rate,
                "sale_inr": round(sale_inr),
                "comm_inr": round(comm_inr),
                "cost_inr": round(cost_inr),
                "gain_inr": round(sale_inr - comm_inr - cost_inr),
                "note": note,
                "fcy_to_usd": 1.0 if s.currency in ("USD", "EUR") else FX_TO_USD_YE2025.get(s.currency, 1.0),
            }
        )
    return rows


def _write_formula_cg(wb, title, subtitle, rows):
    from openpyxl.styles import Font, PatternFill

    ws = wb.create_sheet(title)
    ws["A1"] = subtitle
    ws.merge_cells("A1:U1")
    ws["A2"] = "RECONCILIATION SHEET (formula-driven). Gain/(Loss) INR = Sale INR − Comm INR − Cost INR. Holding >730 → LTCG."
    ws.merge_cells("A2:U2")
    ws["A3"] = "F=E-D | G=IF(F>730,\"LTCG\",\"STCG\") | L=H*K | M=I*K | Q=L*N | R=J*P | S=M*O | T=Q-R-S"
    ws.merge_cells("A3:U3")
    headers = [
        "Symbol", "CCY", "Qty", "Acquisition Date", "Sale Date", "Holding Days", "Type",
        "Sale Proceeds (FCY)", "Cost Basis (FCY)", "Comm (USD)", "FCY→USD Rate",
        "Sale Proceeds (USD)", "Cost (USD)", "Sale FX Rate", "Cost FX Rate", "Comm USD/INR Rate",
        "Sale (INR)", "Comm (INR)", "Cost (INR)", "Gain/(Loss) INR", "Notes",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header(ws, 4, len(headers))
    fill = PatternFill("solid", fgColor="FFF2CC")
    first = 5
    for i, r in enumerate(rows):
        rr = first + i
        usd_sale = rule115_usd(r["sell_date"])
        if r["currency"] == "EUR":
            fcy_to_usd, sale_fx, cost_fx = 1.0, r["sale_rate"], r["cost_rate"]
        else:
            fcy_to_usd = 1.0 if r["currency"] == "USD" else r.get("fcy_to_usd", FX_TO_USD_YE2025.get(r["currency"], 1.0))
            sale_fx = r["sale_rate"] if r["currency"] == "USD" else rule115_usd(r["sell_date"])
            cost_fx = r["cost_rate"] if r["currency"] == "USD" else rule115_usd(r["acq_date"])
        ws.cell(rr, 1, r["symbol"])
        ws.cell(rr, 2, r["currency"])
        ws.cell(rr, 3, r["qty"])
        ws.cell(rr, 4, r["acq_date"])
        ws.cell(rr, 5, r["sell_date"])
        ws.cell(rr, 6, f"=E{rr}-D{rr}")
        ws.cell(rr, 7, f'=IF(F{rr}>730,"LTCG","STCG")')
        ws.cell(rr, 8, round(r["proceeds_local"], 6))
        ws.cell(rr, 9, round(r["cost_local"], 6))
        ws.cell(rr, 10, r["comm_usd"])
        ws.cell(rr, 11, fcy_to_usd)
        ws.cell(rr, 12, f"=H{rr}*K{rr}")
        ws.cell(rr, 13, f"=I{rr}*K{rr}")
        ws.cell(rr, 14, sale_fx)
        ws.cell(rr, 15, cost_fx)
        ws.cell(rr, 16, usd_sale)
        ws.cell(rr, 17, f"=L{rr}*N{rr}")
        ws.cell(rr, 18, f"=J{rr}*P{rr}")
        ws.cell(rr, 19, f"=M{rr}*O{rr}")
        ws.cell(rr, 20, f"=Q{rr}-R{rr}-S{rr}")
        ws.cell(rr, 21, r.get("note") or "")
        for col in (8, 9, 10, 11, 14, 15, 16):
            ws.cell(rr, col).fill = fill
    last = first + len(rows) - 1 if rows else first - 1
    tot = last + 2
    if rows:
        ws.cell(tot, 7, "TOTAL STCG (INR)")
        ws.cell(tot, 20, f'=SUMIF(G{first}:G{last},"STCG",T{first}:T{last})')
        ws.cell(tot + 1, 7, "TOTAL LTCG (INR)")
        ws.cell(tot + 1, 20, f'=SUMIF(G{first}:G{last},"LTCG",T{first}:T{last})')
        ws.cell(tot + 2, 7, "TOTAL CG (INR)")
        ws.cell(tot + 2, 20, f"=T{tot}+T{tot+1}")
        for r in range(tot, tot + 3):
            ws.cell(r, 7).font = Font(bold=True)
            ws.cell(r, 20).font = Font(bold=True)
    autosize(ws)


def _write_os_fy(wb, title, subtitle, dividends, interest):
    ws = wb.create_sheet(title)
    ws["A1"] = subtitle
    ws.append(["Date", "Description", "Currency", "Amount (FCY)", "Amount (USD)", "USD/INR (Rule 115)", "Amount (INR)"])
    style_header(ws, 2, 7)
    int_usd = int_inr = 0.0
    ws.append(["— INTEREST —"])
    for d in interest:
        usd = d["amount"] if d["currency"] == "USD" else to_usd(d["amount"], d["currency"])
        rate = rule115_usd(d["date"])
        inr = usd * rate
        int_usd += usd
        int_inr += inr
        ws.append([d["date"], d["description"], d["currency"], d["amount"], round(usd, 4), rate, round(inr)])
    ws.append([None, "TOTAL INTEREST", None, None, round(int_usd, 2), None, round(int_inr)])
    ws.append([])
    div_usd = div_inr = 0.0
    ws.append(["— DIVIDENDS —"])
    for d in dividends:
        usd = d["amount"] if d["currency"] == "USD" else to_usd(d["amount"], d["currency"])
        rate = rule115_usd(d["date"])
        inr = usd * rate
        div_usd += usd
        div_inr += inr
        ws.append([d["date"], d["description"], d["currency"], d["amount"], round(usd, 4), rate, round(inr)])
    ws.append([None, "TOTAL DIVIDENDS", None, None, round(div_usd, 2), None, round(div_inr)])
    autosize(ws)
    return div_usd, div_inr, int_usd, int_inr


def append_fy2026_27(out: Path, seeds: list[dict]) -> dict:
    from openpyxl import load_workbook

    _, sN = parse_ibkr(FISCAL_FY2627)
    # Inbound GOOG costs controlled by seeds — do not also convert Transfer Ins
    orders = merge_orders(
        extract_orders(sN),
        transfers_to_out_orders(extract_transfers(sN)),
        seeds,
    )
    for o in orders:
        o["symbol"] = normalize_symbol(o["symbol"])
    splits = merge_splits(extract_stock_splits(sN))
    lots_open = build_lots_as_of(orders, splits, date(2026, 3, 31))
    # Only in-window orders — opening lots already include pre-FY seeds (avoid double-apply)
    orders_fy = [o for o in orders if o["date"] >= date(2026, 4, 1)]
    sales, _ = build_fifo_sales(lots_open, orders_fy, date(2026, 4, 1), date(2026, 9, 3), splits=splits)
    cg_rows = _sales_to_cg_rows(sales)
    stcg = sum(r["gain_inr"] for r in cg_rows if r["type"] == "STCG")
    ltcg = sum(r["gain_inr"] for r in cg_rows if r["type"] == "LTCG")

    divs = [d for d in extract_dividends(sN) if date(2026, 4, 1) <= d["date"] <= date(2026, 9, 3)]
    ints = [d for d in extract_interest(sN) if date(2026, 4, 1) <= d["date"] <= date(2026, 9, 3)]
    wht = [w for w in extract_withholding(sN) if date(2026, 4, 1) <= w["date"] <= date(2026, 9, 3)]

    wb = load_workbook(out)
    for name in ("Capital Gains FY2026-27", "Schedule OS - FY2026-27"):
        if name in wb.sheetnames:
            del wb[name]
    _write_formula_cg(
        wb,
        "Capital Gains FY2026-27",
        "Capital gains — FY 2026-27 YTD (01-Apr-2026 to 03-Sep-2026). Partial year. Internal Outs excluded (code I).",
        cg_rows,
    )
    div_usd, div_inr, int_usd, int_inr = _write_os_fy(
        wb,
        "Schedule OS - FY2026-27",
        "Interest & Dividend — FY 2026-27 YTD (01-Apr-2026 to 03-Sep-2026). Partial — full year needs statement to 31-Mar-2027.",
        divs,
        ints,
    )
    if "FTC - Withholding Tax" in wb.sheetnames:
        wsw = wb["FTC - Withholding Tax"]
        for w in wht:
            usd = w["amount"] if w["currency"] == "USD" else to_usd(w["amount"], w["currency"])
            rate = rule115_usd(w["date"])
            wsw.append(["FY2026-27 YTD", w["date"], w["description"], w["currency"], w["amount"], round(usd, 4), rate, round(usd * rate)])
    if "Notes for CA" in wb.sheetnames:
        wn = wb["Notes for CA"]
        wn.append([])
        wn.append(["FY 2026-27 YTD (01-Apr-2026 to 03-Sep-2026) — partial statement"])
        wn.append(["Dividends FY2026-27 YTD", f"USD {div_usd:,.2f} / INR {div_inr:,.0f}"])
        wn.append(["Interest FY2026-27 YTD", f"USD {int_usd:,.2f} / INR {int_inr:,.0f}"])
        wn.append(["STCG FY2026-27 YTD (INR)", f"{stcg:,.0f}"])
        wn.append(["LTCG FY2026-27 YTD (INR)", f"{ltcg:,.0f}"])
        wn.append(["Related Joint", "Internal GOOG transfers with U24577010 (Sofia Anjum Joint) — not taxable CG"])
    if "CG Summary by Symbol" in wb.sheetnames:
        wss = wb["CG Summary by Symbol"]
        wss.append([])
        wss.append(["— FY 2026-27 YTD (to 03-Sep-2026) —"])
        wss.append(["Symbol", "STCG INR", "LTCG INR", "Total Gain/(Loss) INR", "Sale Proceeds USD", "Cost USD"])
        by = {}
        for r in cg_rows:
            e = by.setdefault(r["symbol"], {"stcg": 0, "ltcg": 0, "sale": 0.0, "cost": 0.0})
            if r["type"] == "STCG":
                e["stcg"] += r["gain_inr"]
            else:
                e["ltcg"] += r["gain_inr"]
            e["sale"] += r["proceeds_usd"]
            e["cost"] += r["cost_usd"]
        for sym in sorted(by):
            e = by[sym]
            wss.append([sym, e["stcg"], e["ltcg"], e["stcg"] + e["ltcg"], round(e["sale"], 2), round(e["cost"], 2)])
    wb.save(out)
    print(f"FY2026-27 YTD Individual: CG lots={len(cg_rows)} STCG={stcg:,.0f} LTCG={ltcg:,.0f} Div USD={div_usd:,.2f} INR={div_inr:,.0f}")
    return {"stcg": stcg, "ltcg": ltcg, "div_usd": div_usd, "div_inr": div_inr}


NOTES = [
    (
        "Assessee",
        "Mohammad Waseem — IBKR U25039405 (Individual) — Advisor: MATCAP WEALTH ADVISORS PRIVATE LIMITED",
    ),
    (
        "Statement coverage",
        "Only FY 2026-27 YTD Activity Statement furnished (01-Apr-2026–03-Sep-2026). "
        "FY2025-26 stub; CY2025 Annual synthesized (GOOG × 1 at YE). Full FY2026-27 needs statement to 31-Mar-2027.",
    ),
    (
        "GOOG opening",
        "Prior qty 1 at 01-Apr-2026 seeded 31-Dec-2024 @ USD 98.82. FOP In × 124 on 11-May-2026 @ USD 99.21/sh "
        "(backsolved from IBKR sell Basis). Confirm original buy dates at prior broker for LTCG.",
    ),
    (
        "Related Joint U24577010",
        "Internal Out/In GOOG with Mohammad Waseem & Sofia Anjum Joint — excluded from taxable CG (code I).",
    ),
    (
        "Capital gains FY2026-27 YTD",
        "Taxable: GOOG × 25 sold 01-Jul-2026. Internal Outs not taxable. See Capital Gains FY2026-27 sheet.",
    ),
]


def run():
    seeds = seed_orders()
    synthesize_annual(ANNUAL, seeds)
    main(
        inception=INCEPTION,
        annual=ANNUAL,
        fiscal=FISCAL,
        out=OUT,
        account_open_date="on or before 2024-12-31 (GOOG × 1 held at FY2026-27 start); FOP 11-May-2026",
        assessee_label="Mohammad Waseem",
        notes_extra=NOTES,
        seed_orders=[s for s in seeds if s["date"] <= date(2025, 12, 31)],
    )
    append_fy2026_27(OUT, seeds)


if __name__ == "__main__":
    run()
